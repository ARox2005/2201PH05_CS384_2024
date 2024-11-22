import streamlit as st
import openpyxl as op
from datetime import datetime
from io import BytesIO
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, Border, Side

def scaleWorkbook(inputWorkbook):
    def getColumnLetter(index):
        column = ""
        while index > 0:
            index, remainder = divmod(index - 1, 26)
            column = chr(65 + remainder) + column
        return column

    sheet = inputWorkbook.active
    rowCount = sheet.max_row
    columnCount = sheet.max_column
    studentCount = rowCount - 1

    dataRows = [
        [cell for cell in row] 
        for row in sheet.iter_rows(min_row=2, values_only=True)
    ]
    dataRows.sort(key=lambda row: row[2], reverse=True)

    outputWorkbook = op.Workbook()
    outputSheet = outputWorkbook.active

    gradeData = {
        "AA": {"min": "", "max": ""},
        "AB": {"min": "", "max": ""},
        "BB": {"min": "", "max": ""},
        "BC": {"min": "", "max": ""},
        "CC": {"min": "", "max": ""},
        "CD": {"min": "", "max": ""},
        "DD": {"min": "", "max": ""},
        "F": {"min": "", "max": ""},
    }

    outputSheet.cell(17, 1, "Roll")
    outputSheet.cell(17, 2, "Name")
    outputSheet.cell(17, 3, "Total")
    outputSheet.cell(17, 4, "Grade")
    outputSheet.cell(17, 5, "Scaled")

    for idx, row in enumerate(dataRows, start=18):
        for colIndex, value in enumerate(row):
            outputSheet.cell(idx, colIndex + 1, value)
            if colIndex == 3:
                gradeData[value]["min"] = f"C{idx}"
                if not gradeData[value]["max"]:
                    gradeData[value]["max"] = f"C{idx}"

    dateToday = datetime.now().strftime("%b-%d")

    outputSheet["A1"] = "Subject Code"
    outputSheet["A2"] = "Grade"
    grades = ["AA", "AB", "BB", "BC", "CC", "CD", "DD", "F"]
    gradeBoundaries = [91, 81, 71, 61, 51, 41, 31, 0]

    for i, grade in enumerate(grades, start=3):
        outputSheet[f"A{i}"] = grade
        outputSheet[f"B{i}"] = gradeBoundaries[i - 3]
        outputSheet[f"D{i}"] = gradeData[grade]["min"]
        outputSheet[f"E{i}"] = gradeData[grade]["max"]

    outputSheet["E1"] = dateToday
    outputSheet["F2"] = "Grade"
    outputSheet["G2"] = "Count"

    for i, grade in enumerate(grades, start=3):
        outputSheet[f"G{i}"] = f'=COUNTIF(D18:D{17 + studentCount},"{grade}")'

    img = Image("image.jpg")
    outputSheet.add_image(img, "B12")

    formulaMapping = {
        grade: f"=9*((CX-$D${i})/($E${i}-$D${i}))+$B${i}" for i, grade in enumerate(grades, start=3)
    }

    for rowIdx in range(18, 18 + studentCount):
        outputSheet[f"E{rowIdx}"] = formulaMapping[outputSheet[f"D{rowIdx}"].value].replace("X", str(rowIdx))

    fontStyle = Font(size=15)
    for sheet in outputWorkbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value:
                    cell.font = fontStyle

    colWidths = [15, 37, 17, 10, 17, 12, 8, 4, 15, 12]
    for colIdx, width in enumerate(colWidths, start=1):
        outputSheet.column_dimensions[getColumnLetter(colIdx)].width = width

    borderStyle = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for region in [
        (1, 11, 1, 7),
        (1, 10, 9, 12),
        (17, 17 + studentCount, 1, 5),
    ]:
        for row in outputSheet.iter_rows(*region):
            for cell in row:
                cell.border = borderStyle

    return outputWorkbook


st.title("Excel Grade Processor")
st.write("Upload an Excel file to process and download the updated version.")

uploadedFile = st.file_uploader("Choose an Excel file", type=["xlsx"])

if uploadedFile:
    workbook = op.load_workbook(uploadedFile)
    st.write("Processing your file...")
    outputWorkbook = scaleWorkbook(workbook)

    outputBuffer = BytesIO()
    outputWorkbook.save(outputBuffer)
    outputBuffer.seek(0)

    st.download_button(
        label="Download Processed Excel",
        data=outputBuffer,
        file_name="ProcessedGrades.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    st.success("Processing complete!")
