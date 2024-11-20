import csv
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from datetime import datetime

def is_time_in_range(start_time: str, end_time: str, given_time: str) -> bool:
    time_format = '%H:%M:%S'
    start = datetime.strptime(start_time, time_format).time()
    end = datetime.strptime(end_time, time_format).time()
    time = datetime.strptime(given_time, time_format).time()

    if start <= time <= end:
        return True
    else:
        return False

def set_fill_color(cell, status):
    if status >= 3:
        cell.fill = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')  # Red
    elif status == 1:
        cell.fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')  # Yellow
    elif status == 2:
        cell.fill = PatternFill(start_color='00FF00', end_color='00FF00', fill_type='solid')  # Green

with open('stud_list.txt', 'r') as file:
    student_list = [line.strip() for line in file]

lecture_dates = ["06/08/2024", "13/08/2024", "20/08/2024", "27/08/2024", "03/09/2024", "10/09/2024", "17/09/2024", "01/10/2024"]
attendance_dict = {student: {date: 0 for date in lecture_dates} for student in student_list}

attendance_count = {}

with open('input_attendance.csv', mode='r') as csvfile:
    csv_reader = csv.DictReader(csvfile)
    for row in csv_reader:
        student = row['Roll']
        date = row['Timestamp'][:10]
        time = row['Timestamp'][11:]

        if student in student_list and date in lecture_dates and is_time_in_range('18:00:00', '20:00:00',time)==True:
            if student not in attendance_count:
                attendance_count[student] = {}
            if date not in attendance_count[student]:
                attendance_count[student][date] = 0
            attendance_count[student][date] += 1

for student in attendance_count:
    for date in attendance_count[student]:
        count = attendance_count[student][date]
        attendance_dict[student][date] = count
        # if count == 2:
        #     attendance_dict[student][date] = 2
        # elif count == 1:
        #     attendance_dict[student][date] = 1
        # else:
        #     attendance_dict[student][date] = 0

# print(attendance_dict)

wb = Workbook()
ws = wb.active

ws.append(['Student ID'] + lecture_dates)

for student, attendance in attendance_dict.items():
    row = [student] + [attendance[date] for date in lecture_dates]
    ws.append(row)

for col_idx, date in enumerate(lecture_dates, start=2):
    all_absent = all(attendance_dict[student][date] == 0 for student in student_list)

    if all_absent:
        continue

    for row_idx in range(2, len(student_list) + 2):
        cell = ws.cell(row=row_idx, column=col_idx)
        set_fill_color(cell, cell.value)

wb.save('output_excel.xlsx')

print("Attendance report generated successfully as 'output_excel.xlsx'")
